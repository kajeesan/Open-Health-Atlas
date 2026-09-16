"""Export: any table (optionally a date range) to CSV/XLSX, a multi-table
selective bundle (ZIP of CSVs / one XLSX workbook), and a full-DB download.
Strictly read-only — table names are validated against the live schema (never
interpolated blindly), the date column is derived from the schema (never from
user input), date bounds are bound as `?` params, and the full-DB copy is a
consistent sqlite backup snapshot, not a raw file read mid-write.
"""
import csv
import io
import tempfile
import zipfile
from datetime import date

from flask import (Blueprint, Response, after_this_request, jsonify,
                   request, send_file)

from app import db_read

bp = Blueprint("export", __name__, url_prefix="/api/export")

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _defang(v):
    """Formula-injection guard (CSV *and* XLSX): spreadsheet apps execute text
    cells starting with = + - @ as formulas — a hostile string logged into a
    notes field (chat, Hevy titles) could exfiltrate data the moment the export
    opens in Excel. Prefixing a single quote is Excel's own text escape; the
    value survives visibly, it just can't execute. Numbers pass untouched."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _defang_row(row):
    return {k: _defang(v) for k, v in row.items()}


def _tables():
    """Live table + view names (excludes sqlite internals). This IS the export
    whitelist — the single source of truth for which names are exportable and
    what the UI offers; nothing is hardcoded."""
    rows = db_read.query(
        "SELECT name FROM sqlite_master WHERE type IN ('table','view')"
        " AND name NOT LIKE 'sqlite_%' ORDER BY name")
    return [r["name"] for r in rows]


def _date_col(name):
    """The column a date-range filter applies to, or None. Derived from the
    live schema — NEVER from user input. Every *logged* table uses `date`;
    config tables (recipes, *_products, routines, training_schedule…) have
    none and are exported whole. Caller must have whitelisted `name` first."""
    return "date" if "date" in db_read.columns(name) else None


def _fetch(name, dfrom, dto):
    """(headers, rows) for a whitelisted table, optionally date-filtered.
    `name` is validated against the live schema by the caller; the date column
    is schema-derived (constant `date`), and dfrom/dto are bound as params —
    no user input is ever interpolated into the SQL."""
    dcol = _date_col(name)
    clauses, params = [], []
    if dcol and dfrom:
        clauses.append(f"{dcol} >= ?")  # dcol is the constant 'date', schema-derived
        params.append(dfrom)
    if dcol and dto:
        clauses.append(f"{dcol} <= ?")
        params.append(dto)
    where = (" WHERE " + " AND ".join(clauses)) if clauses else ""
    # name is whitelisted against the live schema -> safe to interpolate
    rows = db_read.query(f"SELECT * FROM '{name}'{where} ORDER BY 1", params)
    headers = list(rows[0].keys()) if rows else db_read.columns(name)
    return headers, rows


def _csv_bytes(headers, rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=headers)
    w.writeheader()
    w.writerows(_defang_row(r) for r in rows)
    return buf.getvalue()


def _xlsx_fill(ws, headers, rows):
    """Write headers + defanged rows into an openpyxl worksheet."""
    ws.append(headers)
    for r in rows:
        # openpyxl treats a "=..." string as a live formula — defang here too
        ws.append([_defang(r.get(h)) for h in headers])


def _send_xlsx(wb, download_name):
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, download_name=download_name, as_attachment=True,
                     mimetype=_XLSX_MIME)


def _sheet_title(name, used):
    """A valid, unique Excel sheet title (≤31 chars). Table names here are
    short sqlite identifiers, but truncation could in principle collide, so we
    de-dupe defensively."""
    base = name[:31]
    title, i = base, 1
    while title.lower() in used:
        suffix = f"~{i}"
        title = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(title.lower())
    return title


def _guard():
    if not db_read.available():
        return jsonify(error="health database not available"), 503
    return None


def _iso_or_none(s):
    """Parse an ISO date; return the date, or None if empty/invalid so the
    caller can 400. Distinguishes 'absent' (arg was empty) via the empty check
    before calling."""
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


@bp.get("/tables")
def tables():
    if (resp := _guard()) is not None:
        return resp
    names = _tables()
    # `dated` = tables a date range can filter; the UI flags the rest as
    # "no date column — full table" so the range control is honest.
    return jsonify(tables=names, dated=[n for n in names if _date_col(n)])


@bp.get("/table")
def table():
    if (resp := _guard()) is not None:
        return resp
    name = request.args.get("name", "")
    if name not in _tables():
        return jsonify(error="unknown table"), 400
    fmt = request.args.get("format", "csv")
    if fmt not in ("csv", "xlsx"):
        return jsonify(error="format must be csv or xlsx"), 400

    dfrom = None
    days = request.args.get("days")
    if days and days.isdigit():
        from app import canon
        dfrom = canon.days_ago_iso(int(days))
    headers, rows = _fetch(name, dfrom, None)

    if fmt == "csv":
        return Response(_csv_bytes(headers, rows), mimetype="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = name[:31]
    _xlsx_fill(ws, headers, rows)
    return _send_xlsx(wb, f"{name}.xlsx")


@bp.get("/bundle")
def bundle():
    """Selective export: N whitelisted tables + an optional ISO date range →
    a ZIP of per-table CSVs (format=csv) or one workbook with a sheet per table
    (format=xlsx). A single table returns a plain .csv/.xlsx (no wrapper)."""
    if (resp := _guard()) is not None:
        return resp

    fmt = request.args.get("format", "csv")
    if fmt not in ("csv", "xlsx"):
        return jsonify(error="format must be csv or xlsx"), 400

    live = set(_tables())
    names, seen = [], set()
    for n in request.args.get("tables", "").split(","):
        n = n.strip()
        if not n or n in seen:
            continue
        if n not in live:
            return jsonify(error=f"unknown table: {n}"), 400
        seen.add(n)
        names.append(n)
    if not names:
        return jsonify(error="select at least one table"), 400

    dfrom = dto = None
    raw_from = request.args.get("from", "").strip()
    raw_to = request.args.get("to", "").strip()
    if raw_from:
        parsed_from = _iso_or_none(raw_from)
        if parsed_from is None:
            return jsonify(error="from must be an ISO date (YYYY-MM-DD)"), 400
        dfrom = parsed_from
    if raw_to:
        parsed_to = _iso_or_none(raw_to)
        if parsed_to is None:
            return jsonify(error="to must be an ISO date (YYYY-MM-DD)"), 400
        dto = parsed_to
    if dfrom and dto and dfrom > dto:
        return jsonify(error="from must be on or before to"), 400
    # `date.fromisoformat` accepts non-canonical ISO forms on Python 3.11+
    # ("20260315", "2026-W11-1") — canonicalize to YYYY-MM-DD before binding
    # into `date >= ?`/`date <= ?` against the stored text column.
    dfrom = dfrom.isoformat() if dfrom else None
    dto = dto.isoformat() if dto else None

    # Single table: plain file, subsuming the quick single-table export 1:1.
    if len(names) == 1:
        headers, rows = _fetch(names[0], dfrom, dto)
        if fmt == "csv":
            return Response(_csv_bytes(headers, rows), mimetype="text/csv",
                            headers={"Content-Disposition": f'attachment; filename="{names[0]}.csv"'})
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = names[0][:31]
        _xlsx_fill(ws, headers, rows)
        return _send_xlsx(wb, f"{names[0]}.xlsx")

    # Multi-table bundle.
    if fmt == "csv":
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for n in names:
                headers, rows = _fetch(n, dfrom, dto)
                z.writestr(f"{n}.csv", _csv_bytes(headers, rows))
        buf.seek(0)
        return send_file(buf, download_name="hermes-export.zip",
                         as_attachment=True, mimetype="application/zip")

    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for n in names:
        ws = wb.create_sheet(title=_sheet_title(n, used))
        headers, rows = _fetch(n, dfrom, dto)
        _xlsx_fill(ws, headers, rows)
    return _send_xlsx(wb, "hermes-export.xlsx")


@bp.get("/database")
def database():
    """Full DB as a consistent snapshot (sqlite backup API, read-only source)."""
    if (resp := _guard()) is not None:
        return resp
    import os
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    db_read.snapshot(tmp.name)

    @after_this_request
    def _cleanup(resp):
        # Never leave a full copy of the health DB lingering in the temp dir.
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
        return resp

    return send_file(tmp.name, download_name="health_metrics.db",
                     as_attachment=True, mimetype="application/octet-stream")
